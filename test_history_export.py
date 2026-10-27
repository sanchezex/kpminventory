import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from history_export import (
    HISTORY_HEADERS,
    INVENTORY_HEADERS,
    export_inventory_png,
    export_inventory_xlsx,
    export_history_png,
    export_history_xlsx,
)


class HistoryExportTests(unittest.TestCase):
    def test_exports_filtered_records_as_excel_and_png(self) -> None:
        record = {
            "movement_date": "2026-04-15T00:00:00+00:00",
            "purchase_date": "2026-04-15",
            "serial_number": "SIM-001",
            "variety_name": "SAFARICOM SIM",
            "from_location": "Samar",
            "to_location": "Proto",
            "destination_company": "Fuel Rod",
            "vehicle_registration": "Tata KCP256C",
            "moved_by": "A. Handler",
            "item_condition": "Good",
            "notes": "Installed and tested",
        }

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "history.xlsx"
            image_path = Path(directory) / "history.png"
            export_history_xlsx(workbook_path, [record])
            export_history_png(image_path, [record])

            sheet = load_workbook(workbook_path, read_only=True).active
            self.assertEqual(tuple(next(sheet.iter_rows(values_only=True))), HISTORY_HEADERS)
            values = tuple(next(sheet.iter_rows(min_row=2, values_only=True)))
            self.assertEqual(values[2], "SIM-001")
            self.assertEqual(values[7], "Tata KCP256C")
            with Image.open(image_path) as image:
                self.assertEqual(image.format, "PNG")
                self.assertGreater(image.width, 1000)
                self.assertGreater(image.height, 100)

    def test_exports_current_inventory_as_excel_and_png(self) -> None:
        record = {
            "serial_number": "SIM-001",
            "variety_name": "SAFARICOM SIM",
            "purchase_date": "2026-04-15T00:00:00+00:00",
            "purchase_cost": 125.5,
            "item_condition": "Good",
            "to_location": "Samar",
            "destination_company": "Internal Store",
        }

        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "inventory.xlsx"
            image_path = Path(directory) / "inventory.png"
            export_inventory_xlsx(workbook_path, [record])
            export_inventory_png(image_path, [record])

            sheet = load_workbook(workbook_path, read_only=True).active
            self.assertEqual(
                tuple(next(sheet.iter_rows(values_only=True))), INVENTORY_HEADERS
            )
            values = tuple(next(sheet.iter_rows(min_row=2, values_only=True)))
            self.assertEqual(values[0], "SIM-001")
            self.assertEqual(values[2], "2026-04-15")
            self.assertEqual(values[3], 125.5)
            self.assertEqual(values[5], "Samar")
            with Image.open(image_path) as image:
                self.assertEqual(image.format, "PNG")
                self.assertGreater(image.width, 900)
                self.assertGreater(image.height, 80)


if __name__ == "__main__":
    unittest.main()

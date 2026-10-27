"""History export helpers for Excel workbooks and PNG reports."""

from __future__ import annotations

import textwrap
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont


HISTORY_HEADERS = (
    "Movement date",
    "Purchase date",
    "Device / serial",
    "Product",
    "From location",
    "To location",
    "Company",
    "Vehicle",
    "Handler",
    "Condition",
    "Notes",
)
HISTORY_FIELDS = (
    "movement_date",
    "purchase_date",
    "serial_number",
    "variety_name",
    "from_location",
    "to_location",
    "destination_company",
    "vehicle_registration",
    "moved_by",
    "item_condition",
    "notes",
)
INVENTORY_HEADERS = (
    "Device / serial",
    "Product",
    "Purchase date",
    "Purchase cost",
    "Condition",
    "Current location",
    "Company",
)
INVENTORY_FIELDS = (
    "serial_number",
    "variety_name",
    "purchase_date",
    "purchase_cost",
    "item_condition",
    "to_location",
    "destination_company",
)


def _export_values(
    row: dict[str, object], fields: tuple[str, ...]
) -> list[str | int | float]:
    values = []
    for field in fields:
        value = row.get(field, "")
        if field == "purchase_cost" and isinstance(value, (int, float)):
            values.append(value)
            continue
        text = str(value) if value is not None else ""
        if field == "purchase_date":
            text = text[:10]
        values.append(text)
    return values


def _export_xlsx(
    file_path: str | Path,
    rows: list[dict[str, object]],
    headers: tuple[str, ...],
    fields: tuple[str, ...],
    sheet_name: str,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(_export_values(row, fields))

    header_fill = PatternFill("solid", fgColor="17324D")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(
            60,
            max(12, max(len(str(cell.value or "")) for cell in column) + 2),
        )
        sheet.column_dimensions[letter].width = width
    workbook.save(file_path)


def _export_png(
    file_path: str | Path,
    rows: list[dict[str, object]],
    title: str,
    headers: tuple[str, ...],
    fields: tuple[str, ...],
    column_widths: tuple[int, ...],
) -> None:
    font = ImageFont.load_default()
    padding = 8
    line_height = 18
    wrapped_rows: list[list[list[str]]] = []

    def wrap_cell(text: str, width: int) -> list[str]:
        character_width = max(8, (width - padding * 2) // 8)
        return textwrap.wrap(text, width=character_width) or [""]

    wrapped_rows.append(
        [
            wrap_cell(header, width)
            for header, width in zip(headers, column_widths)
        ]
    )
    for row in rows:
        wrapped_rows.append(
            [
                wrap_cell(str(value), width)
                for value, width in zip(_export_values(row, fields), column_widths)
            ]
        )

    title_height = 54
    row_heights = [
        max(32, max(len(cell_lines) for cell_lines in row) * line_height + padding * 2)
        for row in wrapped_rows
    ]
    image = Image.new(
        "RGB",
        (sum(column_widths), title_height + sum(row_heights)),
        "#F2F6FA",
    )
    draw = ImageDraw.Draw(image)
    draw.text((padding, 12), title, fill="#17324D", font=font)

    y = title_height
    for row_index, (row, height) in enumerate(zip(wrapped_rows, row_heights)):
        fill = "#17324D" if row_index == 0 else (
            "#FFFFFF" if row_index % 2 else "#E8EEF4"
        )
        draw.rectangle((0, y, image.width, y + height), fill=fill)
        x = 0
        for cell_lines, width in zip(row, column_widths):
            draw.rectangle(
                (x, y, x + width, y + height), outline="#CBD5DF", width=1
            )
            text_color = "#FFFFFF" if row_index == 0 else "#263746"
            for line_index, line in enumerate(cell_lines):
                draw.text(
                    (x + padding, y + padding + line_index * line_height),
                    line,
                    fill=text_color,
                    font=font,
                )
            x += width
        y += height
    image.save(file_path, format="PNG")


def export_history_xlsx(
    file_path: str | Path, rows: list[dict[str, object]]
) -> None:
    _export_xlsx(
        file_path, rows, HISTORY_HEADERS, HISTORY_FIELDS, "Asset history"
    )


def export_history_png(
    file_path: str | Path, rows: list[dict[str, object]]
) -> None:
    _export_png(
        file_path,
        rows,
        "TrackRight asset history",
        HISTORY_HEADERS,
        HISTORY_FIELDS,
        (150, 105, 145, 140, 135, 135, 145, 145, 115, 110, 250),
    )


def export_inventory_xlsx(
    file_path: str | Path, rows: list[dict[str, object]]
) -> None:
    _export_xlsx(
        file_path, rows, INVENTORY_HEADERS, INVENTORY_FIELDS, "Current inventory"
    )


def export_inventory_png(
    file_path: str | Path, rows: list[dict[str, object]]
) -> None:
    _export_png(
        file_path,
        rows,
        "TrackRight current inventory",
        INVENTORY_HEADERS,
        INVENTORY_FIELDS,
        (145, 150, 115, 115, 110, 145, 145),
    )
